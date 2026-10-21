"""Excel download of a bill run: one sheet per plant in the agreed 17 columns,
then Summary, Review (lines not billed, with the reason) and PO Rates."""
import io
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from ..db import one, rows
from .po_parser import price_on

F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
HF = PatternFill("solid", fgColor="D9E1F2")
YF = PatternFill("solid", fgColor="FFF2CC")
OF = PatternFill("solid", fgColor="F8CBAD")   # Existing Rate = billed rate (PO export had only the new rate)
BF = PatternFill("solid", fgColor="DDEBF7")   # Existing Rate taken from another PO of the same part
BD = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
COLS = ["Voucher Number", "Date", "Party Name", "Order Number", "Order Due Date", "Item Name", "", "Item HSN",
        "Item rate of GST", "Billed Quantity", "Rate", "Existing Rate", " Reversed  Rate", "Diff.Rate", "Basic Value",
        "Grin No", "Grin date"]
W = [12, 11, 34, 12, 12, 14, 32, 10, 8, 9, 9, 10, 10, 9, 12, 13, 11]


def _n(v):
    s = str(v or "")
    return int(s) if s.isdigit() and len(s) < 16 else v


def _d(iso):
    try:
        return date.fromisoformat(iso[:10]) if iso else None
    except ValueError:
        return iso


def _header(ws, cols, widths, row=1):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row, c, h)
        x.font, x.fill, x.border = FB, HF, BD
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(c)].width = w


def _lines(ws, data, start, extra=(), overrides=()):
    r = start
    for x in data:
        vals = [x["voucher"], _d(x["inv_date"]), x["party"], _n(x["order_no"]), _d(x["order_due"]), x["part"],
                x["descr"], _n(x["hsn"]), x["gst"], x["qty"], x["rate"], x["old_rate"], x["new_rate"],
                f'=IF(OR(L{r}="",M{r}=""),"",ROUND(M{r}-L{r},2))', f'=IF(N{r}="","",ROUND(N{r}*J{r},2))',
                _n(x["grn_no"]), _d(x["grn_date"])] + [x[e] for e in extra]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v).font = F
        ws.cell(r, 2).number_format = ws.cell(r, 5).number_format = "DD-MMM-YY"
        ws.cell(r, 17).number_format = "DD/MM/YYYY"
        for c in (11, 12, 13, 14, 15):
            ws.cell(r, c).number_format = "#,##0.00"
        ws.cell(r, 16).number_format = "0"
        if x["part"] in overrides:
            ws.cell(r, 13).fill = YF
        src = x.get("old_src") or "PO"
        if src == "billed rate":
            ws.cell(r, 12).fill = OF
        elif src != "PO":
            ws.cell(r, 12).fill = BF
        r += 1
    return r


def build(run_id, diff_basis="po"):
    run = one("SELECT * FROM bill_runs WHERE id=?", (run_id,))
    params = json.loads(run["params"])
    ovr = {o["part"] for o in params.get("overrides", [])}
    data = rows("SELECT * FROM run_lines WHERE run_id=? ORDER BY plant, part, inv_date, voucher", (run_id,))
    bill = [x for x in data if x["cat"] == "Bill"]
    plants = sorted({x["plant"] for x in bill}, key=lambda p: -sum(x["value"] for x in bill if x["plant"] == p))
    wb = Workbook()
    wb.remove(wb.active)
    tot = {}
    for p in plants:
        ws = wb.create_sheet(p[:31])
        _header(ws, COLS, W)
        r = _lines(ws, [x for x in bill if x["plant"] == p], 2, overrides=ovr)
        if diff_basis == "billed":   # Diff = Reversed - billed Rate
            for i in range(2, r):
                ws.cell(i, 14).value = f'=IF(M{i}="","",ROUND(M{i}-K{i},2))'
        end = r - 1
        ws.cell(r, 1, "TOTAL").font = FB
        for c, col in ((10, "J"), (15, "O")):
            x = ws.cell(r, c, f"=SUM({col}2:{col}{end})")
            x.font, x.fill, x.number_format = FB, HF, "#,##0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:Q{end}"
        tot[p] = (end, r)

    pend = [x for x in data if x["cat"] in PENDING_CATS]
    pend.sort(key=lambda x: (x["cat"], x["inv_date"] or "", x["voucher"], x["part"]))
    pend_plants = sorted({x["plant"] for x in pend}, key=lambda p: (plants.index(p) if p in plants else 99, p))
    for p in pend_plants:
        _pending_sheet(wb, f"{p} GRN PENDING", [x for x in pend if x["plant"] == p], ovr)

    s = wb.create_sheet("Summary")
    s["A1"] = f"Supplementary Bill - TAFE price revision {params['old_date']} -> {params['new_date']} (plant-wise)"
    s["A1"].font = Font(name="Arial", size=12, bold=True)
    _header(s, ["Plant", "Lines", "Total Qty", "Basic Value (Rs)", "GST (Rs)", "Total incl. GST (Rs)"],
            [14, 10, 12, 16, 14, 18], row=3)
    r = 4
    for p in plants:
        end, t = tot[p]
        q = f"'{p}'"
        vals = [p, f"=COUNTA({q}!A2:A{end})", f"={q}!J{t}", f"={q}!O{t}",
                f"=SUMPRODUCT({q}!O2:O{end},{q}!I2:I{end})/100", f"=D{r}+E{r}"]
        for c, v in enumerate(vals, 1):
            x = s.cell(r, c, v)
            x.font, x.border = F, BD
            if c >= 3:
                x.number_format = "#,##0.00"
        r += 1
    s.cell(r, 1, "TOTAL").font = FB
    for c in range(2, 7):
        x = s.cell(r, c, f"=SUM({L(c)}4:{L(c)}{r - 1})")
        x.font, x.fill, x.number_format = FB, HF, ("#,##0.00" if c > 2 else "0")
    r += 2
    s.cell(r, 1, "Bill value by PO").font = FB
    r += 1
    _header(s, ["Plant", "PO No", "Lines", "Basic Value (Rs)"], [14, 10, 12, 16], row=r)
    r += 1
    by_po = {}
    for x in bill:
        k = (x["plant"], x["po_no"])
        n, v = by_po.get(k, (0, 0.0))
        by_po[k] = (n + 1, v + x["value"])
    for (p, po), (n, v) in sorted(by_po.items()):
        for c, val in enumerate([p, _n(po), n, round(v, 2)], 1):
            x = s.cell(r, c, val)
            x.font, x.border = F, BD
            if c == 4:
                x.number_format = "#,##0.00"
        r += 1
    r += 1
    if pend_plants:
        s.cell(r, 1, "Pending - billed once TAFE gives the GRIN / revised rate (see GRN PENDING sheets)").font = FB
        r += 1
        _header(s, ["Plant", "Pending for", "Lines", "Basic Value (Rs)"], [14, 30, 12, 16], row=r)
        r += 1
        for p in pend_plants:
            for cat, label in (("GRN pending", "GRIN not received"), ("Rate pending", "Revised rate not received")):
                xs = [x for x in pend if x["plant"] == p and x["cat"] == cat]
                if xs:
                    v = round(sum(x["value"] or 0 for x in xs), 2)
                    for c, val in enumerate([p, label, len(xs), v if v else "rate not known"], 1):
                        x = s.cell(r, c, val)
                        x.font, x.border = F, BD
                        if c == 4 and v:
                            x.number_format = "#,##0.00"
                    r += 1
        r += 1
    s.cell(r, 1, "Not in the bill (see Review sheet)").font = FB
    r += 1
    _header(s, ["Plant", "Category", "Lines"], [14, 30, 12], row=r)
    r += 1
    cnt = {}
    for x in data:
        if x["cat"] != "Bill":
            cnt[(x["plant"], x["cat"])] = cnt.get((x["plant"], x["cat"]), 0) + 1
    for (p, c_), n in sorted(cnt.items()):
        for c, v in enumerate([p, c_, n], 1):
            s.cell(r, c, v).font = F
            s.cell(r, c).border = BD
        r += 1
    r += 1
    notes = ["Notes",
             f"Existing Rate = PO price valid on {params['old_date']}; Reversed Rate = price from {params['new_date']} "
             "(Gross + Pkg.Charges + Freight from the PO).",
             "Diff.Rate = Reversed - Existing; Basic Value = Diff.Rate x Billed Quantity." if diff_basis != "billed"
             else "Diff.Rate = Reversed - billed Rate; Basic Value = Diff.Rate x Billed Quantity.",
             "Fixed Reversed Rates (Settings): " + (", ".join(f"{o['part']} = {o['revised']}" for o in params.get("overrides", [])) or "none")
             + " - highlighted yellow.",
             "Grin No / date taken from the TAFE GRIN files, checked against SDEW's GRN sheet.",
             "Existing Rate in blue = taken from another PO of the same part with the same Reversed Rate "
             "(TAFE PO export shows only the new rate).",
             "Existing Rate in orange = PO export shows only the new rate and no other PO has the part: "
             "Existing Rate = the last rate billed before the revision (highest billed rate below the new rate). Lines: "
             + (", ".join(f"{p} {sum(1 for x in bill if x['plant'] == p and x.get('old_src') == 'billed rate')} lines "
                          f"Rs {sum(x['value'] for x in bill if x['plant'] == p and x.get('old_src') == 'billed rate'):,.2f}"
                          for p in plants) or "none") + "."]
    for n in notes:
        s.cell(r, 1, n).font = FB if n == "Notes" else F
        r += 1
    s.column_dimensions["B"].width = 30

    rv = wb.create_sheet("Review")
    _header(rv, COLS + ["Plant", "Category", "Reason", "GRIN check", "Remarks"], W + [10, 16, 34, 14, 40])
    rest = sorted((x for x in data if x["cat"] != "Bill"), key=lambda x: (x["plant"], x["cat"], x["reason"], x["part"]))
    r = _lines(rv, rest, 2, extra=("plant", "cat", "reason", "status", "notes"), overrides=ovr)
    rv.freeze_panes = "A2"
    rv.auto_filter.ref = f"A1:V{max(r - 1, 1)}"

    pr = wb.create_sheet("PO Rates")
    _header(pr, ["PO No", "Plant", "Item", "Part No", "Description", "Existing Rate", "Reversed Rate", "Diff"],
            [12, 10, 6, 15, 30, 13, 13, 9])
    r = 2
    ovr_map = {o["part"]: o["revised"] for o in params.get("overrides", [])}
    for it in rows("SELECT i.*, h.plant FROM po_items i LEFT JOIN po_headers h USING(po_no) ORDER BY po_no, item"):
        b = json.loads(it["blocks"])
        o, n = price_on(b, params["old_date"]), price_on(b, params["new_date"])
        if it["part"] in ovr_map and o is not None:
            n = ovr_map[it["part"]]
        for c, v in enumerate([_n(it["po_no"]), it["plant"], it["item"], it["part"], it["descr"], o, n,
                               f'=IF(OR(F{r}="",G{r}=""),"",G{r}-F{r})'], 1):
            pr.cell(r, c, v).font = F
        for c in (6, 7, 8):
            pr.cell(r, c).number_format = "#,##0.00"
        r += 1
    pr.freeze_panes = "A2"
    pr.auto_filter.ref = f"A1:H{max(r - 1, 1)}"

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PENDING_CATS = ("GRN pending", "Rate pending")


def _pending_sheet(wb, title, xs, ovr):
    for x in xs:
        x["pending_for"] = "GRIN" if x["cat"] == "GRN pending" else ("Rate + GRIN" if not x["grn_no"] else "Rate")
    ws = wb.create_sheet(title[:31])
    _header(ws, COLS + ["Pending for", "Reason", "Remarks"], W + [13, 40, 44])
    end = _lines(ws, xs, 2, extra=("pending_for", "reason", "notes"), overrides=ovr) - 1
    t = end + 1
    ws.cell(t, 1, "TOTAL").font = FB
    for c, col in ((10, "J"), (15, "O")):
        x = ws.cell(t, c, f"=SUM({col}2:{col}{end})")
        x.font, x.fill, x.number_format = FB, HF, "#,##0.00"
    for i in range(2, end + 1):
        if ws.cell(i, 18).value != "GRIN":   # revised rate not received - yellow
            ws.cell(i, 18).fill = YF
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:T{max(end, 1)}"
    return ws


def build_pending(run_id):
    """Pending list: lines that will be billed once TAFE gives the GRIN (GRN pending)
    or the revised rate (Rate pending), one sheet per plant, in the bill columns."""
    run = one("SELECT * FROM bill_runs WHERE id=?", (run_id,))
    params = json.loads(run["params"])
    ovr = {o["part"] for o in params.get("overrides", [])}
    data = rows("SELECT * FROM run_lines WHERE run_id=? AND cat IN (?,?) ORDER BY plant, cat, inv_date, voucher, part",
                (run_id, *PENDING_CATS))
    order = {"DBR": 0, "MDU": 1}
    plants = sorted({x["plant"] for x in data}, key=lambda p: (order.get(p, 9), p))
    wb = Workbook()
    wb.remove(wb.active)
    s = wb.create_sheet("Summary")
    s["A1"] = "Supplementary Bill - pending list (GRIN not received / revised rate not received)"
    s["A1"].font = Font(name="Arial", size=12, bold=True)
    _header(s, ["Plant", "Pending for", "Lines", "Qty", "Basic Value (Rs)"], [14, 26, 10, 12, 18], row=3)
    r = 4
    for p in plants:
        for cat in PENDING_CATS:
            xs = [x for x in data if x["plant"] == p and x["cat"] == cat]
            if not xs:
                continue
            v = round(sum(x["value"] or 0 for x in xs), 2)
            label = "GRIN pending" if cat == "GRN pending" else "Revised rate pending"
            for c, val in enumerate([p, label, len(xs), sum(x["qty"] or 0 for x in xs), v if v else "rate not known"], 1):
                x = s.cell(r, c, val)
                x.font, x.border = F, BD
                if c == 5 and v:
                    x.number_format = "#,##0.00"
            r += 1
    r += 1
    rp = sorted({x["reason"].split(":")[0] for x in data if x["cat"] == "Rate pending"})
    for n in ["Notes",
              "GRIN pending: value is claimable once TAFE posts the GRIN for the invoice.",
              "Rate pending: TAFE has not given the revised rate yet (" + (", ".join(rp) or "none")
              + ") - Reversed Rate and value left blank."]:
        s.cell(r, 1, n).font = FB if n == "Notes" else F
        r += 1
    for p in plants:
        _pending_sheet(wb, p, [x for x in data if x["plant"] == p], ovr)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
