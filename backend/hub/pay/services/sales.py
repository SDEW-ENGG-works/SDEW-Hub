"""Sales marking: GRN NO / GRN DATE / GRN QTY in the 'TAFE SALES DETAILS' file from the hub's GRIN lines.

The file (one row per invoice line):
  INVOICE NO | INVOICE DATE | Month | AMOUNT | plant | PART NO | DESCRIPTION | QTY | GRN QTY | GRN NO | GRN DATE | HSN
GRIN lines come from the shared grin_lines table (Fetch GRIN from TAFE Channel / GRIN Details upload).

Matching, per invoice (invoice no cleaned, so 9552/26/27 = 9552/26-27):
  1. same part number (cleaned: spaces, '-', MO->M0)                    -> best
  2. one part number starts with the other (1827A039016 ~ 1827A0390169) -> next
  3. invoice has one sales line and one GRIN line, or the same quantity -> "part differs" note
  Each GRIN line is used once while unused ones are left (tool sets: 3 rows -> 3 GRIN lines).
  Several GRIN lots for one row: the one with the same quantity, else the latest.
Written:  GRN NO (number), GRN DATE (dd/mm/yyyy text, as typed by hand), GRN QTY = accepted quantity.
  Inspection not finished (accepted 0 and rejected 0) -> GRN QTY left blank.
  A GRN NO / GRN QTY already in the file is never changed unless the line is ticked.
"""
import io
import json
import os
import re
from collections import OrderedDict
from datetime import date, datetime

from ...grin.services.norm import inv_key
from ...http import ApiError
from ...supp.services.norm import part_key, to_iso
from ...supp.services.readers import sheets
from ..db import get_db, get_setting, one, rows, set_setting

CATS = ["Fill GRN", "Add GRN QTY", "Short / rejected", "Inspection open", "GRN no differs", "GRN QTY differs",
        "Already complete", "GRIN pending", "Not in GRIN data", "Not a GRIN party"]
MARK_CATS = {"Fill GRN", "Add GRN QTY", "Short / rejected", "Inspection open"}
LOCKED = {"Already complete", "GRIN pending", "Not in GRIN data", "Not a GRIN party"}

FIELDS = {
    "inv": ("invoice no", "invoice number", "inv no"),
    "date": ("invoice date", "inv date"),
    "party": ("plant", "party", "customer", "party name"),
    "part": ("part no", "part number", "item code"),
    "desc": ("description", "item name"),
    "qty": ("qty", "quantity", "billed qty", "inv qty"),
    "gqty": ("grn qty", "grin qty", "accepted qty", "grn quantity"),
    "gno": ("grn no", "grin no", "grn number"),
    "gdate": ("grn date", "grin date"),
}


# ----------------------------------------------------------------------------- reading the file
def _header(row):
    cells = [re.sub(r"\s+", " ", str(v).strip().lower()) if v is not None else "" for v in row]
    cols = {}
    for f, names in FIELDS.items():
        for n in names:
            if n in cells:
                cols[f] = cells.index(n)
                break
    return cols if {"inv", "part", "gno"} <= set(cols) else None


def locate(path):
    """(sheet, header row, columns) of the sales sheet. With several such sheets: the one with a GRN QTY
    column, then the one with most rows."""
    best = None
    for name, it in sheets(path):
        cols, hdr, n = None, 0, 0
        for i, row in enumerate(it, start=1):
            if cols is None:
                if i > 15:
                    break
                cols = _header(row)
                hdr = i
                continue
            if row and row[cols["inv"]] not in (None, ""):
                n += 1
        if cols:
            score = ("gqty" in cols, n)
            if best is None or score > best[0]:
                best = (score, name, hdr, cols)
    return best[1:] if best else None


def _num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_sales(path):
    loc = locate(path)
    if not loc:
        raise ValueError("No 'INVOICE NO' + 'PART NO' + 'GRN NO' header found")
    sheet, hdr, cols = loc
    out = []
    for name, it in sheets(path):
        if name != sheet:
            continue
        for i, row in enumerate(it, start=1):
            if i <= hdr:
                continue
            g = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
            inv = g("inv")
            if inv in (None, "") or not str(inv).strip():
                continue
            out.append(dict(row=i, inv=_text(inv), key=inv_key(inv), date=to_iso(g("date")), party=_text(g("party")),
                            part=_text(g("part")), pkey=part_key(g("part")), desc=_text(g("desc")), qty=_num(g("qty")),
                            gno=_text(g("gno")), gdate=_text(g("gdate")) if not isinstance(g("gdate"), (date, datetime))
                            else to_iso(g("gdate")), gqty=_num(g("gqty"))))
        break
    return dict(sheet=sheet, header_row=hdr, cols=cols, rows=out)


def save_sales_file(name, path, kind="Sales file"):
    from .ingest import _add_file
    s = read_sales(path)
    info = dict(sheet=s["sheet"], lines=len(s["rows"]), invoices=len({r["key"] for r in s["rows"]}),
                with_grn=sum(1 for r in s["rows"] if r["gno"]), with_qty=sum(1 for r in s["rows"] if r["gqty"] is not None),
                has_qty_col="gqty" in s["cols"],
                last=max((r["date"] for r in s["rows"] if r["date"]), default=""))
    fid = _add_file(kind, name, path, len(s["rows"]), info)
    set_setting("sales_file", fid)
    msg = (f"{info['lines']:,} lines ({info['invoices']:,} invoices) on sheet '{s['sheet']}' · "
           f"{info['with_grn']:,} with GRN NO · {info['with_qty']:,} with GRN QTY")
    if not info["has_qty_col"]:
        msg += " · no 'GRN QTY' column - add it to fill the accepted quantity"
    return dict(kind=kind, file_id=fid, message=msg, info=info)


# ----------------------------------------------------------------------------- matching
def grin_parties():
    """Tally party names (cleaned) of the plants that have a TAFE GRIN plant code (Supp Bill > Settings > Plants)."""
    from ...grin.services.defaults import plants
    return {re.sub(r"\s+", " ", p["party"].strip().lower()): p["code"] for p in plants() if (p.get("grin_plant") or "").strip()}


def _pscore(a, b):
    if not a or not b:
        return 0
    if a == b:
        return 3
    if min(len(a), len(b)) >= 8 and (a.startswith(b) or b.startswith(a)):
        return 2
    return 0


def _open(g):
    return not (g["accepted_qty"] or 0) and not (g["rejected_qty"] or 0)


def _same_no(a, b):
    a, b = _text(a), _text(b)
    if a.endswith(".0"):
        a = a[:-2]
    return a == b


def run_sales():
    from .ingest import file_path
    fid = get_setting("sales_file")
    path = file_path(fid) if fid else None
    if not path:
        raise ApiError("Upload the TAFE SALES DETAILS file first.")
    sales = read_sales(path)
    has_qty = "gqty" in sales["cols"]
    grins = rows("SELECT grin_no, line_no, grin_date, invoice_no, part_no, part_desc, grin_qty, accepted_qty, "
                 "rejected_qty, plant FROM grin_lines")
    if not grins:
        raise ApiError("No GRIN data in the hub yet - press Fetch GRIN.")
    upto = max(g["grin_date"] for g in grins if g["grin_date"])
    by_inv = {}
    for g in grins:
        g["pkey"] = part_key(g["part_no"])
        by_inv.setdefault(inv_key(g["invoice_no"]), []).append(g)
    parties = grin_parties()
    today = date.today()

    inv_rows = OrderedDict()
    for r in sales["rows"]:
        inv_rows.setdefault(r["key"], []).append(r)

    out = []
    used_grins = set()
    for key, srows in inv_rows.items():
        glines = by_inv.get(key, [])
        used = set()
        picks = {}
        # pass 1: part number (exact, then prefix); prefer GRIN lines not used yet by this invoice
        for r in srows:
            cands = [(g, _pscore(r["pkey"], g["pkey"])) for g in glines]
            cands = [(g, s) for g, s in cands if s]
            if not cands:
                continue
            top = max(s for _, s in cands)
            best = [g for g, s in cands if s == top]
            fresh = [g for g in best if id(g) not in used] or best
            same_qty = [g for g in fresh if r["qty"] is not None and g["grin_qty"] == r["qty"]]
            pick = (same_qty or sorted(fresh, key=lambda g: (g["grin_date"], g["grin_no"]), reverse=True))[0]
            others = [g for g in best if g is not pick and g["pkey"] == pick["pkey"] and g["grin_no"] != pick["grin_no"]]
            picks[id(r)] = (pick, "prefix" if top == 2 else "", others)
            used.add(id(pick))
        # pass 2: rows still open - one line each side, or the same quantity
        left = [g for g in glines if id(g) not in used]
        for r in srows:
            if id(r) in picks or not left:
                continue
            m = None
            if len(srows) == 1 and len(glines) == 1:
                m = glines[0]
            else:
                q = [g for g in left if r["qty"] is not None and g["grin_qty"] == r["qty"]]
                m = q[0] if len(q) == 1 else None
            if m:
                picks[id(r)] = (m, "part", [])
                used.add(id(m))
                left = [g for g in left if g is not m]
        used_grins |= {(g["grin_no"], g["line_no"]) for g in glines if id(g) in used}

        for r in srows:
            rec = dict(row=r["row"], inv=r["inv"], inv_key=key, inv_date=r["date"], party=r["party"], part=r["part"],
                       qty=r["qty"], sheet_gno=r["gno"], sheet_gdate=r["gdate"], sheet_gqty=r["gqty"],
                       grin_no="", grin_date="", grin_qty=None, acc_qty=None, rej_qty=None, grin_part="")
            notes = []
            pk = picks.get(id(r))
            if not pk:
                party_ok = re.sub(r"\s+", " ", r["party"].lower()) in parties
                if glines:
                    cat = "Not in GRIN data"
                    notes.append(f"GRIN has this invoice for part {', '.join(sorted({g['part_no'] for g in glines}))}")
                elif r["gno"]:
                    cat = "Not in GRIN data"
                    notes.append(f"File has GRN {r['gno']}, GRIN data (up to {upto}) has no line for this invoice")
                elif not party_ok:
                    cat = "Not a GRIN party"
                else:
                    cat = "GRIN pending"
                    if r["date"]:
                        days = (today - date.fromisoformat(r["date"])).days
                        notes.append(f"{days} days since invoice")
            else:
                g, how, others = pk
                rec.update(grin_no=g["grin_no"], grin_date=g["grin_date"], grin_qty=g["grin_qty"],
                           acc_qty=g["accepted_qty"], rej_qty=g["rejected_qty"], grin_part=g["part_no"])
                if how == "prefix":
                    notes.append(f"Part in GRIN: {g['part_no']}")
                elif how == "part":
                    notes.append(f"Part differs - GRIN has {g['part_no']}, matched on {'quantity' if len(glines) > 1 else 'invoice'}")
                if others:
                    notes.append("Other GRIN lots for this part: " + ", ".join(f"{o['grin_no']} ({o['grin_qty']:g})" for o in others))
                inspection = _open(g)
                short = not inspection and r["qty"] is not None and (g["accepted_qty"] or 0) < r["qty"]
                if r["gno"] and not _same_no(r["gno"], g["grin_no"]):
                    cat = "GRN no differs"
                    notes.insert(0, f"File has {r['gno']}, TAFE GRIN is {g['grin_no']} ({g['grin_date']}). Tick to correct it")
                elif has_qty and r["gqty"] is not None and not inspection and r["gqty"] != g["accepted_qty"]:
                    cat = "GRN QTY differs"
                    notes.insert(0, f"File has GRN QTY {r['gqty']:g}, TAFE accepted {g['accepted_qty']:g}. Tick to correct it")
                elif short:
                    cat = "Short / rejected"
                    notes.insert(0, f"Accepted {g['accepted_qty']:g} of {r['qty']:g}"
                                    + (f", rejected {g['rejected_qty']:g}" if g["rejected_qty"] else ""))
                elif inspection:
                    cat = "Inspection open"
                    notes.insert(0, "TAFE inspection not finished - GRN QTY left blank")
                elif not r["gno"]:
                    cat = "Fill GRN"
                elif has_qty and r["gqty"] is None:
                    cat = "Add GRN QTY"
                else:
                    cat = "Already complete"
                # nothing to write when the row already has everything that can be filled
                if cat in ("Short / rejected", "Inspection open") and r["gno"] and (
                        inspection or not has_qty or r["gqty"] is not None):
                    nothing = True
                else:
                    nothing = False
                rec["_nothing"] = nothing
            rec["cat"] = cat
            rec["action"] = "mark" if cat in MARK_CATS and not rec.pop("_nothing", False) else "skip"
            rec.pop("_nothing", None)
            rec["note"] = "; ".join(notes)
            out.append(rec)

    out.sort(key=lambda x: (CATS.index(x["cat"]), x["row"]))
    unmatched = [g for g in grins if (g["grin_no"], g["line_no"]) not in used_grins
                 and inv_key(g["invoice_no"]) not in inv_rows]
    last_sales = max((r["date"] for r in sales["rows"] if r["date"]), default="")
    first_sales = min((r["date"] for r in sales["rows"] if r["date"]), default="")
    unmatched = [g for g in unmatched if first_sales <= (g["grin_date"] or "")]
    summary = _summary(out, has_qty, upto, unmatched, last_sales)
    # all lines of the check go into ONE compressed field of ps_runs (~0.3 MB instead of ~12 MB of rows),
    # so the check is quick also when the database is Turso and the office internet is slow
    con = get_db()
    cur = con.execute("INSERT INTO ps_runs(created_at, sales_file, summary, lines) VALUES(?,?,?,?)",
                      (datetime.now().isoformat(timespec="seconds"), fid, json.dumps(summary), _pack_lines(out)))
    rid = cur.lastrowid
    con.execute("DELETE FROM ps_actions WHERE run_id <= ?", (rid - 10,))
    con.execute("DELETE FROM ps_lines")        # rows of checks made before 29-Oct-2026 (no longer used)
    con.execute("DELETE FROM ps_runs WHERE id <= ?", (rid - 10,))
    con.commit()
    return sales_summary(rid)


LINE_COLS = ["cat", "action", "row", "inv", "inv_key", "inv_date", "party", "part", "qty", "sheet_gno", "sheet_gdate",
             "sheet_gqty", "grin_no", "grin_date", "grin_qty", "acc_qty", "rej_qty", "grin_part", "note"]
_cache = {}


def _pack_lines(out):
    import zlib
    return zlib.compress(json.dumps([[o[c] for c in LINE_COLS] for o in out], separators=(",", ":")).encode(), 6)


def base_lines(rid):
    """The lines of a check as stored (default tick per line). Kept in memory per check."""
    import zlib
    if rid not in _cache:
        r = one("SELECT lines FROM ps_runs WHERE id=?", (rid,))
        if not r or not r["lines"]:
            raise ApiError("This check was made by an older version - press Check GRN again.")
        data = json.loads(zlib.decompress(bytes(r["lines"])).decode())
        _cache.clear()
        _cache[rid] = [dict(zip(LINE_COLS, v), seq=i) for i, v in enumerate(data)]
    return _cache[rid]


def run_lines(rid):
    """Lines with the ticks the user changed: one line (seq) beats its whole group (cat), which beats the default."""
    seq_act, cat_act = {}, {}
    for a in rows("SELECT seq, cat, action FROM ps_actions WHERE run_id=?", (rid,)):
        if a["seq"] is not None and a["seq"] >= 0:
            seq_act[a["seq"]] = a["action"]
        else:
            cat_act[a["cat"]] = a["action"]
    out = []
    for l in base_lines(rid):
        act = l["action"]
        if l["grin_no"] and l["cat"] not in LOCKED and l["cat"] in cat_act:
            act = cat_act[l["cat"]]
        if l["seq"] in seq_act:
            act = seq_act[l["seq"]]
        out.append(dict(l, action=act) if act != l["action"] else l)
    return out


def set_line_action(rid, seq, action):
    con = get_db()
    con.execute("DELETE FROM ps_actions WHERE run_id=? AND seq=?", (rid, seq))
    con.execute("INSERT INTO ps_actions(run_id, seq, cat, action) VALUES(?,?,?,?)", (rid, seq, "", action))
    con.commit()


def set_cat_action(rid, cat, action):
    seqs = [l["seq"] for l in base_lines(rid) if l["cat"] == cat]
    con = get_db()
    for i in range(0, len(seqs), 300):
        part = seqs[i:i + 300]
        con.execute(f"DELETE FROM ps_actions WHERE run_id=? AND seq IN ({','.join('?' * len(part))})", [rid] + part)
    con.execute("DELETE FROM ps_actions WHERE run_id=? AND seq=-1 AND cat=?", (rid, cat))
    con.execute("INSERT INTO ps_actions(run_id, seq, cat, action) VALUES(?,?,?,?)", (rid, -1, cat, action))
    con.commit()


def _summary(out, has_qty, upto, unmatched, last_sales):
    cats = {c: dict(n=0, marking=0) for c in CATS}
    for o in out:
        cats[o["cat"]]["n"] += 1
        cats[o["cat"]]["marking"] += o["action"] == "mark"
    matched = [o for o in out if o["grin_no"]]
    agree = sum(1 for o in matched if o["sheet_gno"] and _same_no(o["sheet_gno"], o["grin_no"]))
    typed = sum(1 for o in matched if o["sheet_gno"])
    return dict(cats=cats, lines=len(out), has_qty_col=has_qty, grin_upto=upto, last_sales=last_sales,
                agree=agree, typed=typed,
                grin_not_in_file=[dict(grin_no=g["grin_no"], grin_date=g["grin_date"], invoice_no=g["invoice_no"],
                                       part_no=g["part_no"], grin_qty=g["grin_qty"], plant=g["plant"])
                                  for g in sorted(unmatched, key=lambda g: g["grin_date"], reverse=True)[:300]],
                grin_not_in_file_n=len(unmatched))


def latest_sales_run():
    r = one("SELECT MAX(id) id FROM ps_runs")
    return r["id"] if r else None


def sales_summary(rid):
    run = one("SELECT id, created_at, sales_file, summary, exported_at, exported_file FROM ps_runs WHERE id=?", (rid,))
    if not run:
        return None
    run["summary"] = json.loads(run["summary"])
    try:
        lines = run_lines(rid)
    except ApiError:
        return None
    for c in run["summary"]["cats"].values():
        c["marking"] = 0
    for l in lines:
        if l["action"] == "mark":
            run["summary"]["cats"][l["cat"]]["marking"] += 1
    run["summary"]["to_mark"] = sum(c["marking"] for c in run["summary"]["cats"].values())
    f = one("SELECT name FROM pm_files WHERE id=?", (run["sales_file"],))
    run["sales_name"] = f["name"] if f else ""
    return run


# ----------------------------------------------------------------------------- writing the file
def dmy(iso_date):
    return f"{iso_date[8:10]}/{iso_date[5:7]}/{iso_date[0:4]}" if iso_date and len(iso_date) >= 10 else (iso_date or "")


def _qty_value(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else v


def marked_sales(rid):
    import openpyxl
    from .ingest import _store, file_path
    run = sales_summary(rid)
    if not run:
        raise ApiError("Run the sales check first.")
    path = file_path(run["sales_file"])
    if not path or not os.path.exists(path):
        raise ApiError("The sales file of this check is missing - upload it again and re-check.")
    sheet, hdr, cols = locate(path)
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    n = 0
    for l in [x for x in run_lines(rid) if x["action"] == "mark" and x["grin_no"]]:
        r = l["row"]
        ws.cell(r, cols["gno"] + 1).value = int(l["grin_no"]) if l["grin_no"].isdigit() else l["grin_no"]
        if "gdate" in cols:
            ws.cell(r, cols["gdate"] + 1).value = dmy(l["grin_date"])
        if "gqty" in cols:
            inspection = not (l["acc_qty"] or 0) and not (l["rej_qty"] or 0)
            ws.cell(r, cols["gqty"] + 1).value = None if inspection else _qty_value(l["acc_qty"])
        n += 1
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    base = os.path.splitext(run["sales_name"] or "TAFE SALES DETAILS.xlsx")[0].split("_GRN_")[0]
    name = f"{base}_GRN_{date.today().strftime('%d%m%Y')}.xlsx"
    saved = save_sales_file(name, _store(name, data), kind="Sales file (marked)")
    con = get_db()
    con.execute("UPDATE ps_runs SET exported_at=?, exported_file=? WHERE id=?",
                (datetime.now().isoformat(timespec="seconds"), saved["file_id"], rid))
    con.commit()
    return data, name, n


def sales_report(rid):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from .export import _table
    run = sales_summary(rid)
    if not run:
        raise ApiError("Run the sales check first.")
    sm = run["summary"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SDEW — Sales marking (GRN)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"File: {run['sales_name']} · GRIN data up to {dmy(sm['grin_upto'])} · checked "
                f"{run['created_at'].replace('T', ' ')} · GRN NO typed by hand agrees with TAFE: {sm['agree']:,} of {sm['typed']:,}")
    shown = [c for c in CATS if sm["cats"][c]["n"]]
    _table(ws, ["Result", "Lines", "Written now"], [[c, sm["cats"][c]["n"], sm["cats"][c]["marking"]] for c in shown],
           [30, 12, 12], start=4)
    ws.freeze_panes = None
    ws2 = wb.create_sheet("Lines")
    data = run_lines(rid)
    _table(ws2, ["Result", "Written", "Row", "Invoice No", "Invoice date", "Party", "Part No", "Qty", "File GRN NO",
                 "File GRN DATE", "File GRN QTY", "TAFE GRIN no", "GRIN date", "GRIN qty", "Accepted", "Rejected",
                 "GRIN part", "Note"],
           [[l["cat"], "Yes" if l["action"] == "mark" else "", l["row"], l["inv"], dmy(l["inv_date"]), l["party"],
             l["part"], l["qty"], l["sheet_gno"], l["sheet_gdate"], l["sheet_gqty"], l["grin_no"], dmy(l["grin_date"]),
             l["grin_qty"], l["acc_qty"], l["rej_qty"], l["grin_part"], l["note"]] for l in data],
           [20, 8, 8, 13, 12, 32, 16, 8, 13, 12, 10, 13, 12, 9, 9, 9, 16, 60])
    fills = {"Fill GRN": "FFEAF6EF", "Add GRN QTY": "FFEAF6EF", "Short / rejected": "FFFBF3E2",
             "Inspection open": "FFFBF3E2", "GRN no differs": "FFFBECEA", "GRN QTY differs": "FFFBECEA",
             "GRIN pending": "FFF2F4F7", "Not in GRIN data": "FFFBECEA"}
    for r, l in enumerate(data, 2):
        if fills.get(l["cat"]):
            ws2.cell(r, 1).fill = PatternFill("solid", fgColor=fills[l["cat"]])
    ws2.auto_filter.ref = f"A1:R{len(data) + 1}"
    if sm["grin_not_in_file"]:
        ws3 = wb.create_sheet("GRIN not in file")
        _table(ws3, ["GRIN no", "GRIN date", "Invoice No", "Part No", "GRIN qty", "Plant"],
               [[g["grin_no"], dmy(g["grin_date"]), g["invoice_no"], g["part_no"], g["grin_qty"], g["plant"]]
                for g in sm["grin_not_in_file"]], [14, 12, 14, 16, 10, 8])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Sales_GRN_Check_{date.today().strftime('%d%m%Y')}.xlsx"
