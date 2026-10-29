"""Excel outputs: the marked register, and a check report."""
import io
import os
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..db import get_db, rows
from ...http import ApiError
from .defaults import load_settings
from .engine import CATS, run_summary
from .ingest import file_path, save_register, _store
from ...grin.services.norm import dotted
from .register import locate


def marked_register(rid):
    """Writes voucher no + date into the register for every ticked line. Returns (bytes, filename, count)."""
    run = run_summary(rid)
    if not run:
        raise ApiError("Run the check first.")
    path = file_path(run["register_file"])
    if not path or not os.path.exists(path):
        raise ApiError("The register file of this check is missing — upload it again and re-check.")
    s = load_settings()
    sheet, hdr, cols = locate(path)
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    mc, dc = cols["mark"] + 1, cols.get("mark_date", cols["mark"] + 1) + 1
    last = max(mc, dc)
    fill = PatternFill("solid", fgColor="FF" + s["fill_color"].lstrip("#").upper()[-6:]) if s.get("fill_rows") else None
    n = 0
    for l in rows("SELECT * FROM pm_run_lines WHERE run_id=? AND action='mark' AND reg_row IS NOT NULL", (rid,)):
        r = l["reg_row"]
        ws.cell(r, mc).value = int(l["pv"]) if l["pv"].isdigit() else l["pv"]
        ws.cell(r, dc).value = dotted(l["pv_date"])
        if fill:      # off by default: the row is coloured by hand after checking
            for c in range(1, last + 1):
                ws.cell(r, c).fill = fill
        n += 1
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    base = os.path.splitext(run["register_name"] or "Register.xlsx")[0]
    base = base.split("_marked_")[0]
    name = f"{base}_marked_{date.today().strftime('%d%m%Y')}.xlsx"
    # the marked file becomes the register for the next check
    saved = save_register(name, _store(name, data), kind="Register (marked)")
    con = get_db()
    con.execute("UPDATE pm_runs SET exported_at=?, exported_file=? WHERE id=?",
                (datetime.now().isoformat(timespec="seconds"), saved["file_id"], rid))
    con.commit()
    return data, name, n


HEAD_FILL = PatternFill("solid", fgColor="FF1E3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFFFF")
THIN = Border(*(Side(style="thin", color="FFD0D5DD"),) * 4)
CAT_FILL = {"To mark": "FFEAF6EF", "Credit + debit note": "FFEAF6EF", "Amount differs": "FFFBF3E2",
            "Marked with another number": "FFFBF3E2", "Debit note only": "FFF2F4F7",
            "Not in register": "FFFBECEA", "Cancelled in register": "FFFBECEA"}


def _table(ws, headers, data, widths, start=1):
    for i, h in enumerate(headers, 1):
        c = ws.cell(start, i, h)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, THIN
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(data, start + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.border = THIN
            if isinstance(v, float):
                c.number_format = "#,##0.00"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(start + 1, 1)


def check_report(rid):
    run = run_summary(rid)
    if not run:
        raise ApiError("Run the check first.")
    sm = run["summary"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "SDEW — Payment check"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Register: {run['register_name']} · checked {run['created_at'].replace('T', ' ')}"
    _table(ws, ["Result", "Invoices", "Marked now", "Credit (Rs.)", "Debit (Rs.)"],
           [[c, sm["cats"][c]["n"], sm["cats"][c]["marking"], sm["cats"][c]["credit"], sm["cats"][c]["debit"]]
            for c in CATS if sm["cats"][c]["n"]], [30, 12, 12, 16, 16], start=4)
    r0 = 6 + len([c for c in CATS if sm["cats"][c]["n"]]) + 1
    ws.cell(r0 - 1, 1, "Payment vouchers").font = Font(bold=True)
    _table(ws, ["Voucher", "Date", "Amount (Rs.)", "Credit (Rs.)", "Debit (Rs.)", "Cash discount", "Credit − Debit", "Tallies"],
           [[v["pv"], dotted(v["pv_date"]), v["amount"], v["credit"], v["debit"], v["discount"], v["net"],
             "Yes" if v["tallies"] else "NO"] for v in sm["vouchers"]], [30, 12, 16, 16, 16, 14, 16, 10], start=r0)
    ws.freeze_panes = None

    ws2 = wb.create_sheet("Invoices")
    data = rows("SELECT * FROM pm_run_lines WHERE run_id=? ORDER BY seq", (rid,))
    _table(ws2, ["Result", "Marked", "Invoice No", "Register row", "Invoice date", "Party", "Invoice amount",
                 "Voucher", "Voucher date", "Credit (Rs.)", "Debit (Rs.)", "Doc types", "Register mark", "Note"],
           [[l["cat"], "Yes" if l["action"] == "mark" else "", l["inv"], l["reg_row"], dotted(l["reg_date"]),
             l["reg_party"], l["reg_amount"], l["pv"], dotted(l["pv_date"]), l["credit"], l["debit"], l["doc_types"],
             (f"{l['reg_mark']} {l['reg_mark_date']}".strip()), l["note"]] for l in data],
           [24, 8, 14, 10, 12, 34, 14, 12, 12, 14, 12, 10, 24, 60])
    for r, l in enumerate(data, 2):
        f = CAT_FILL.get(l["cat"])
        if f:
            ws2.cell(r, 1).fill = PatternFill("solid", fgColor=f)
    ws2.auto_filter.ref = f"A1:N{len(data) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"Payment_Check_{date.today().strftime('%d%m%Y')}.xlsx"
