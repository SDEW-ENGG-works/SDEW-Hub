import io
import urllib.parse
from datetime import date, datetime

from .. import Response, route
from ..db import get_db, rows
from ...http import ApiError
from ..services import status as st
from ..services.norm import inv_key


def _as_of(req):
    a = req.arg("as_of").strip()
    if a:
        try:
            date.fromisoformat(a)
        except ValueError:
            raise ApiError("as_of must be yyyy-mm-dd")
    return a or None


@route("GET", r"/api/status")
def status(req):
    data = st.build(_as_of(req))
    out = st.filtered(data, req.arg("status"), req.arg("plant"), req.arg("q"), req.arg("buyer"))
    off = int(req.arg("offset", "0") or 0)
    lim = min(int(req.arg("limit", "200") or 200), 2000)
    return dict(as_of=data["as_of"], counts=data["counts"], order=data["order"], dates=data["dates"],
                plants=data["plants"], settings=data["settings"], untracked=data["untracked"],
                total=len(out), value=round(sum(r["amount"] or 0 for r in out), 2), rows=out[off:off + lim])


@route("GET", r"/api/invoice")
def invoice(req):
    k = inv_key(urllib.parse.unquote(req.arg("no")))
    data = st.build(_as_of(req))
    row = next((r for r in data["rows"] if r["inv_key"] == k), None)
    grin = [g for g in rows("SELECT * FROM grin_lines ORDER BY grin_date") if inv_key(g["invoice_no"]) == k]
    pay = rows("SELECT * FROM gp_pay_lines WHERE inv_key=? ORDER BY pv_date", (k,))
    reg = rows("SELECT * FROM gp_register WHERE inv_key=?", (k,))
    mails = rows("SELECT m.id, m.created_at, m.sent_at, m.to_addr, m.status FROM gp_mail_invoices mi "
                 "JOIN gp_mails m ON m.id=mi.mail_id WHERE mi.inv_key=? ORDER BY m.id", (k,))
    tally = rows("SELECT * FROM gp_tally_invoices WHERE inv_key=?", (k,))
    if not (row or grin or pay or reg or tally):
        raise ApiError("Invoice not found", 404)
    return dict(inv_key=k, row=row, grin=grin, pay=pay, register=reg, mails=mails, tally=tally)


@route("PUT", r"/api/note")
def note(req):
    b = req.json
    k = inv_key(b.get("inv_key") or "")
    if not k:
        raise ApiError("inv_key missing")
    con = get_db()
    con.execute("INSERT INTO gp_notes(inv_key, note, hold, updated_at) VALUES(?,?,?,?) ON CONFLICT(inv_key) DO UPDATE "
                "SET note=excluded.note, hold=excluded.hold, updated_at=excluded.updated_at",
                (k, (b.get("note") or "").strip()[:500], 1 if b.get("hold") else 0,
                 datetime.now().isoformat(timespec="seconds")))
    con.commit()
    st.invalidate()
    return dict(ok=True)


@route("GET", r"/api/export")
def export(req):
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    data = st.build(_as_of(req))
    out = st.filtered(data, req.arg("status"), req.arg("plant"), req.arg("q"), req.arg("buyer"))
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("GRIN & Payments")
    head = ["Status", "Invoice No", "Invoice Date", "Plant", "Party", "PO No", "Parts", "Qty", "Value", "Days",
            "GRIN No", "GRIN Date", "GRIN Qty", "Rejected", "Payment Voucher", "Payment Date", "Paid Amount",
            "Due Date", "Overdue Days", "Buyer", "Reminders", "Last Reminder", "Hold", "Note", "Checks"]
    widths = [20, 13, 12, 10, 34, 13, 40, 8, 12, 6, 14, 12, 9, 9, 15, 12, 12, 12, 9, 18, 9, 12, 6, 30, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    hfill, hfont = PatternFill("solid", fgColor="1E293B"), Font(bold=True, color="FFFFFF")
    hdr = []
    for h in head:
        c = WriteOnlyCell(ws, value=h)
        c.fill, c.font = hfill, hfont
        hdr.append(c)
    ws.append(hdr)
    fills = {k: PatternFill("solid", fgColor=v) for k, v in {
        st.NOT_GRINED: "FECACA", st.SHORT: "FED7AA", st.OVERDUE: "FDE68A", st.AWAIT: "DBEAFE",
        st.WAIT: "E2E8F0", st.DEBIT: "E9D5FF", st.PAID: "BBF7D0"}.items()}
    date_cols, money_cols = {2, 11, 15, 17, 21}, {8, 16}

    def d(v):
        try:
            return datetime.strptime(v[:10], "%Y-%m-%d") if v else None
        except ValueError:
            return v
    for r in out:
        vals = [r["status"], r["inv_no"], d(r["inv_date"]), r["plant"], r["party"], r["order_no"],
                ", ".join(f"{p['part']} x {p['qty']:g}" for p in r["parts"]), r["qty"], r["amount"], r["age"],
                r["grin_nos"], d(r["grin_date"]), r["grin_qty"], r["rejected"], r["pv"], d(r["pv_date"]), r["credit"],
                d(r["due_date"]), r["overdue_days"], r["buyer"], r["reminders"], d(r["last_reminder"]),
                "Yes" if r["hold"] else "", r["note"], "; ".join(r["notes"])]
        cells = []
        for i, v in enumerate(vals):
            if i == 0 or i in date_cols or i in money_cols:
                c = WriteOnlyCell(ws, value=v)
                if i == 0:
                    c.fill = fills.get(r["status"]) or PatternFill()
                elif i in date_cols:
                    c.number_format = "dd-mm-yyyy"
                else:
                    c.number_format = "#,##0.00"
                cells.append(c)
            else:
                cells.append(v)
        ws.append(cells)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(out) + 1}"
    sm = wb.create_sheet("Summary")
    sm.append(["As of", data["as_of"]])
    sm.append(["GRIN data up to", data["dates"]["grin_until"]])
    sm.append(["Payment data up to", data["dates"]["pay_until"]])
    sm.append([])
    sm.append(["Status", "Invoices", "Value"])
    for s_ in data["order"]:
        sm.append([s_, data["counts"][s_]["n"], data["counts"][s_]["value"]])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"GRIN_Payment_Status_{data['as_of']}.xlsx"
    return Response(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
